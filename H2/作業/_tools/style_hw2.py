# -*- coding: utf-8 -*-
"""把 2026-MKT-HW2.docx 排版美化：不改任何文字與表格內容，只調樣式。
- 題目（英文原題）加上灰底區塊，和答案明顯分開
- (一)(二)(三)(四) 做成深藍色標題橫條
- 小節標題（1. 2. 3.）藍色粗體
- 內文中文用微軟正黑體、左右對齊、行距放寬
- 表格深藍表頭 + 白字、隔行淺藍、淡灰框線、表頭跨頁重複
"""
import re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

WD = r"d:\OneDrive\OneDrive - PEGATRON CORPORATION\交大\國際行銷\H2\作業"
SRC = WD + r"\2026-MKT-HW2.docx"

FONT = "微軟正黑體"
C_TITLE = "1F4E79"      # 深藍
C_SECTION = "1F4E79"    # 章節橫條底
C_SUB = "2E75B6"        # 小節藍
C_BODY = "1A1A1A"
C_PROMPT_FILL = "EDF1F5"  # 題目灰底
C_HEAD_FILL = "1F4E79"    # 表頭
C_ZEBRA = "EEF4FA"        # 隔行
C_BORDER = "B7C9DB"

doc = docx.Document(SRC)


# ---------- 底層 XML 小工具 ----------
def rfonts(rpr):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    return rf


def shade(pr, fill):
    old = pr.find(qn("w:shd"))
    if old is not None:
        pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def p_shade(p, fill):
    shade(p._p.get_or_add_pPr(), fill)


def cell_shade(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def p_bottom_border(p, color, sz=8, space=1):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn("w:pBdr"))
    if old is not None:
        pPr.remove(old)
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), str(sz))
    bot.set(qn("w:space"), str(space))
    bot.set(qn("w:color"), color)
    bdr.append(bot)
    pPr.append(bdr)


# 模組層級的 w:pPr 子元素順序表（w:shd / w:pBdr 必須照 schema 擺）
PPR_ORDER = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore",
             "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
             "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
             "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
             "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
             "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
             "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
             "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
             "w:rPr", "w:sectPr", "w:pPrChange"]


def fix_ppr_order(p):
    """w:shd / w:pBdr 用 append 會排在錯的位置，改成依 schema 排序。"""
    pPr = p._p.find(qn("w:pPr"))
    if pPr is None:
        return
    kids = list(pPr)
    def key(el):
        tag = el.tag.split("}")[1]
        name = "w:" + tag
        return PPR_ORDER.index(name) if name in PPR_ORDER else 99
    for el in sorted(kids, key=key):
        pPr.append(el)


def style_run(r, size=None, bold=None, color=None, italic=None):
    if size is not None:
        r.font.size = Pt(size)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if color is not None:
        r.font.color.rgb = RGBColor.from_string(color)
    rfonts(r._element.get_or_add_rPr())


def set_spacing(p, before=None, after=None, line=None, rule=None):
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if rule is not None:
        pf.line_spacing_rule = rule


# ---------- 全域字型與頁面 ----------
normal = doc.styles["Normal"]
normal.font.size = Pt(11.5)
rfonts(normal.element.get_or_add_rPr())

sec = doc.sections[0]
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)


# ---------- 分類 ----------
RE_SECTION = re.compile(r"^\([一二三四五六]\)")
RE_SUB = re.compile(r"^\d+\.\s")


def is_prompt(t):
    return t.startswith("Based on") or t.startswith("Regarding")


def classify(t):
    if not t.strip():
        return "blank"
    if is_prompt(t):
        return "prompt"
    if t.startswith("圖 1"):
        return "caption"
    if RE_SECTION.match(t):
        return "section"
    if RE_SUB.match(t):
        return "sub"
    return "body"


# 標題（第一段）
title = doc.paragraphs[0]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in title.runs:
    style_run(r, size=20, bold=True, color=C_TITLE)
set_spacing(title, before=0, after=16, line=1.15)
p_bottom_border(title, C_TITLE, sz=12, space=6)
fix_ppr_order(title)

counts = {}
for p in doc.paragraphs[1:]:
    kind = classify(p.text)
    counts[kind] = counts.get(kind, 0) + 1

    if kind == "blank":
        # 表格之間的空段落：壓扁，避免表格間距過大
        set_spacing(p, before=0, after=0, line=Pt(6))
        for r in p.runs:
            style_run(r, size=6)
        continue

    if kind == "prompt":
        p_shade(p, C_PROMPT_FILL)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_spacing(p, before=16, after=12, line=1.3)
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.right_indent = Cm(0.25)
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            style_run(r, size=11, bold=True, color="1F3A5F")
        fix_ppr_order(p)

    elif kind == "section":
        p_shade(p, C_SECTION)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_spacing(p, before=16, after=10, line=1.2)
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            style_run(r, size=13.5, bold=True, color="FFFFFF")
        fix_ppr_order(p)

    elif kind == "sub":
        set_spacing(p, before=10, after=4, line=1.3)
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            style_run(r, size=11.5, bold=True, color=C_SUB)
        fix_ppr_order(p)

    elif kind == "caption":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=4, after=14, line=1.15)
        for r in p.runs:
            style_run(r, size=10, bold=True, color="5A6B7B")
        fix_ppr_order(p)

    else:  # body
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_spacing(p, before=0, after=7, line=1.45)
        p.paragraph_format.left_indent = Cm(0.3)
        for r in p.runs:
            style_run(r, size=11.5, color=C_BODY)
        fix_ppr_order(p)


# ---------- 表格 ----------
def set_borders(t):
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge, color, sz in (("top", C_BORDER, "6"), ("left", C_BORDER, "6"),
                            ("bottom", C_BORDER, "6"), ("right", C_BORDER, "6"),
                            ("insideH", C_BORDER, "4"), ("insideV", C_BORDER, "4")):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def long_col(t, j):
    """整欄的文字都很短（<=10 字）才置中，長文字的欄位保持靠左。"""
    return max(len(str(t.cell(i, j).text)) for i in range(len(t.rows))) > 10


for t in doc.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    set_borders(t)

    ncols = len(t.columns)
    center_cols = {j for j in range(ncols) if j > 0 and not long_col(t, j)}

    for i, row in enumerate(t.rows):
        if i == 0:
            # 表頭跨頁重複
            trPr = row._tr.get_or_add_trPr()
            hdr = OxmlElement("w:tblHeader")
            trPr.append(hdr)

        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i == 0:
                cell_shade(cell, C_HEAD_FILL)
            elif i % 2 == 0:
                cell_shade(cell, C_ZEBRA)

            for p in cell.paragraphs:
                set_spacing(p, before=2, after=2, line=1.15)
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if j in center_cols
                               else WD_ALIGN_PARAGRAPH.LEFT)
                if not p.runs and p.text == "":
                    style_run(p.add_run(""), size=10)
                for r in p.runs:
                    style_run(r, size=10,
                              bold=(i == 0),
                              color=("FFFFFF" if i == 0 else C_BODY))

doc.save(SRC)
print("styled ok. para kinds:", counts)
d2 = docx.Document(SRC)
print("paragraphs:", len(d2.paragraphs), "tables:", len(d2.tables),
      "inline_shapes:", len(d2.inline_shapes))
