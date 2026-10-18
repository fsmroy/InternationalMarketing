# 在 Q4 第(二)節「由座標可以看出…」段落後插入 MDS 定位圖 + 圖說
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"d:\OneDrive\OneDrive - PEGATRON CORPORATION\交大\國際行銷\H2\作業\2026-MKT-HW2.docx"
doc = docx.Document(SRC)

anchor = next(p for p in doc.paragraphs if p.text.startswith("由座標可以看出"))

pic_p = doc.paragraphs[-1]          # add_picture 會把圖片塞在最後一段
doc.add_picture(r"_tools\mds_plot.png", width=Inches(5.6))
pic_p = doc.paragraphs[-1]
pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
anchor._p.addnext(pic_p._p)          # 圖片移到 anchor 之後

cap = doc.add_paragraph()            # 圖說，之後移到圖片後面
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.add_run("圖 1　8 款手機在 Dim 1 / Dim 2 定位平面上的分布").bold = True
pic_p._p.addnext(cap._p)

doc.save(SRC)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
print("inline_shapes:", len(doc.inline_shapes))
